# -*- coding: utf-8 -*-
import io,os,zipfile,shutil,glob,json
from lxml import etree
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
def L(e): return etree.QName(e).localname
q=lambda tag:'{http://www.hancom.co.kr/hwpml/2011/paragraph}'+tag
def subtext(e): return ''.join((x.text or '') for x in e.iter() if L(x)=='t')
def celltext(tc): return ' '.join(''.join((x.text or '') for x in p.iter() if L(x)=='t') for p in tc.iter(q('p'))).strip()
def pt(p): return ''.join((x.text or '') for run in p.findall('{*}run') for x in run if L(x)=='t')
def tonum(s):
    s=(s or '').replace(',','').replace('%','').strip()
    try: return float(s)
    except: return None
def load(src):
    tmp='be_%d'%abs(hash(src))
    if os.path.exists(tmp): shutil.rmtree(tmp)
    os.makedirs(tmp)
    with zipfile.ZipFile(src) as z: z.extractall(tmp)
    FP=max(glob.glob(tmp+'/Contents/section*.xml'),key=os.path.getsize)
    return [e for e in etree.parse(FP).iter() if L(e)=='sec'][0]
def title_prev(kids,tb):
    # find paragraph index containing tb
    for i,p in enumerate(kids):
        if tb in p.iter():
            for j in range(i-1,-1,-1):
                if pt(kids[j]).strip(): return pt(kids[j]).strip()
    return ''
def tbl_title(sec,kids,tb):
    c=tb.find(q('caption'))
    return subtext(c).strip() if c is not None else title_prev(kids,tb)
def rows_of(tb):
    res=[]
    for tr in tb.findall(q('tr')):
        res.append([celltext(x) for x in tr.findall(q('tc'))])
    return res

REPORTS={
 '디스플레이':'C:/Users/remoteadmin/Desktop/CNL/CNL보고서_초안/디스플레이_최종.hwpx',
 '반도체':'C:/Users/remoteadmin/Desktop/CNL/CNL보고서_초안/반도체_최종(종합의견추가).hwpx',
 '로봇':'C:/Users/remoteadmin/Desktop/CNL/CNL보고서_초안/로봇산업_최종.hwpx',
 '이차전지':'C:/Users/remoteadmin/Desktop/CNL/CNL보고서_초안/이차전지_최종.hwpx',
}
def analyze(src):
    sec=load(src); kids=[p for p in sec if L(p)=='p']
    na1=None; da1=None; item_cnt=None
    for tb in sec.iter(q('tbl')):
        title=tbl_title(sec,kids,tb)
        rows=rows_of(tb)
        if not rows: continue
        hdr=rows[0]
        # 나-1: 주요국별 논문 수 및 점유율
        if na1 is None and '주요국별' in title and '논문' in title and '점유율' in title and len(hdr)>=3:
            na1=rows
        # 나-2: 핵심 품목별 ... 점유율 -> item count
        if item_cnt is None and ('핵심 품목별' in title) and ('점유율' in title):
            item_cnt=len([r for r in rows[1:] if r and r[0].strip()])
        # 다-1: 연도별 R&D 과제 수 및 정부투자연구비
        if da1 is None and '연도별' in title and '과제' in title and '정부투자연구비' in title and any('과제' in h for h in hdr):
            da1=rows
    # compute paper scale
    paper_sum=None; share_sum=None; countries=[]
    if na1:
        for r in na1[1:]:
            if not r or not r[0].strip(): continue
            nm=r[0].strip()
            if '기타' in nm or '합계' in nm: continue
            countries.append(nm)
            n=tonum(r[1]); s=tonum(r[2]) if len(r)>2 else None
            if n is not None: paper_sum=(paper_sum or 0)+n
            if s is not None: share_sum=(share_sum or 0)+s
    # compute R&D scale
    task_sum=None; inv_sum=None
    if da1:
        hdr=da1[0]
        # find columns for 과제수 and 투자비
        ci_task=next((k for k,h in enumerate(hdr) if '과제' in h),None)
        ci_inv=next((k for k,h in enumerate(hdr) if '정부투자' in h or '연구비' in h),None)
        for r in da1[1:]:
            if not r or not r[0].strip(): continue
            if '합계' in r[0]: continue
            if ci_task is not None and ci_task<len(r):
                v=tonum(r[ci_task]); task_sum=(task_sum or 0)+v if v is not None else task_sum
            if ci_inv is not None and ci_inv<len(r):
                v=tonum(r[ci_inv]); inv_sum=(inv_sum or 0)+v if v is not None else inv_sum
    return dict(countries=countries,paper_sum=paper_sum,share_sum=share_sum,
                item_cnt=item_cnt,task_sum=task_sum,inv_sum=inv_sum)

info={}
for name,src in REPORTS.items():
    info[name]=analyze(src)
    print(name, info[name]['countries'], info[name]['paper_sum'], info[name]['item_cnt'], info[name]['task_sum'], info[name]['inv_sum'])
io.open('overview_info.json','w',encoding='utf-8').write(json.dumps(info,ensure_ascii=False,indent=1))

# ---------- write Excel ----------
def man(n): return '약 %.1f만 편'%(n/10000) if n else '-'
def eok(n): return '약 %s억원'%('{:,}'.format(int(round(n)))) if n else '-'
wb=openpyxl.Workbook(); wb.remove(wb.active)
hdr_fill=PatternFill('solid',fgColor='1F4E79'); hdr_font=Font(color='FFFFFF',bold=True,size=11,name='맑은 고딕')
sub_fill=PatternFill('solid',fgColor='DDEBF7'); lab_font=Font(bold=True,name='맑은 고딕',size=10)
base_font=Font(name='맑은 고딕',size=10)
thin=Side(style='thin',color='BFBFBF'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical='center')
for name,src in REPORTS.items():
    d=info[name]
    ws=wb.create_sheet(name)
    ws.column_dimensions['A'].width=16; ws.column_dimensions['B'].width=72
    rr=[1]
    def title_row(txt):
        r=rr[0]
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
        c=ws.cell(r,1,txt); c.fill=hdr_fill; c.font=hdr_font; c.alignment=Alignment(horizontal='left',vertical='center')
        ws.row_dimensions[r].height=20; rr[0]+=1
    def kv(k,v):
        r=rr[0]
        a=ws.cell(r,1,k); a.fill=sub_fill; a.font=lab_font; a.alignment=wrap; a.border=border
        b=ws.cell(r,2,v); b.font=base_font; b.alignment=wrap; b.border=border
        rr[0]+=1
    countries='·'.join(d['countries']) if d['countries'] else '주요 5개국'
    n_country=len(d['countries']) if d['countries'] else 5
    title_row('< 논문 데이터 개요 >')
    kv('자료원','Web of Science(SCIE) 논문 데이터베이스')
    kv('분석 기간','2022~2026년(최근 5년)')
    kv('분석 대상','%s 분야 SCIE 논문(%s개 핵심 품목)'%(name, d['item_cnt'] or ''))
    kv('대상 국가','주요 %d개국(%s)'%(n_country, countries))
    kv('집계 기준','국가 중복 계수(논문 1건에 여러 국가가 참여하면 각 국가에 각각 계수)')
    kv('분석 규모','주요 %d개국 합계 %s(전체의 약 %.1f%%)'%(n_country, man(d['paper_sum']), d['share_sum'] or 0))
    rr[0]+=1
    title_row('< R&D 데이터 개요 >')
    kv('자료원','NTIS(국가과학기술지식정보서비스) 국가연구개발과제 데이터')
    kv('분석 기간','2022~2026년(최근 5년)')
    kv('분석 대상','%s 관련 국가연구개발(R&D) 과제(%s개 핵심 품목)'%(name, d['item_cnt'] or ''))
    kv('집계 기준','과제 수 및 정부투자연구비(억원)')
    kv('분석 규모','5년 누적 총 %s개 과제, 정부투자연구비 %s'%('{:,}'.format(int(d['task_sum'])) if d['task_sum'] else '-', eok(d['inv_sum'])))
OUTX='데이터_개요_산업별.xlsx'
wb.save(OUTX)
print('saved', OUTX)
